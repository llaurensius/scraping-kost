import io
import json
from typing import List, Dict, Any

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


class ExportService:
    """
    Generator ekspor dataset kost ke file format CSV dan Excel (.xlsx)
    sesuai spesifikasi PRD & TECHNICAL_SPEC.
    """

    COLUMN_MAPPING = [
        ("id", "ID Kost"),
        ("name", "Nama Kost"),
        ("gender_type", "Tipe Gender"),
        ("price_monthly", "Harga Bulanan (Rp)"),
        ("vfm_score", "Skor VFM (0-100)"),
        ("vfm_category", "Kategori VFM"),
        ("motorcycle_distance_km", "Jarak Motor (km)"),
        ("motorcycle_duration_min", "Waktu Motor (mnt)"),
        ("driving_distance_km", "Jarak Mobil (km)"),
        ("driving_duration_min", "Waktu Mobil (mnt)"),
        ("walking_distance_km", "Jarak Jalan (km)"),
        ("walking_duration_min", "Waktu Jalan (mnt)"),
        ("straight_distance_km", "Garis Lurus (km)"),
        ("room_available", "Sisa Kamar"),
        ("room_size", "Ukuran Kamar"),
        ("rating", "Rating"),
        ("review_count", "Jumlah Review"),
        ("facilities_room", "Fasilitas Kamar"),
        ("owner_phone", "No WhatsApp/Telp"),
        ("full_address", "Alamat Lengkap"),
        ("district", "Kecamatan"),
        ("city", "Kota"),
        ("is_detail_enriched", "Status Enrich"),
        ("url_detail", "Link Mamikos"),
        ("url_gmaps_route", "Rute Google Maps"),
        ("scraped_at", "Waktu Scrape")
    ]

    @classmethod
    def _prepare_flat_data(cls, items: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        rows = []
        for item in items:
            row = {}
            for key, header_name in cls.COLUMN_MAPPING:
                val = item.get(key)
                if key == "facilities_room":
                    if isinstance(val, list):
                        val = ", ".join(str(f) for f in val)
                    elif isinstance(val, str) and val.startswith("["):
                        try:
                            parsed = json.loads(val)
                            val = ", ".join(str(f) for f in parsed)
                        except Exception:
                            pass
                elif key == "is_detail_enriched":
                    val = "Lengkap (Enriched)" if val == 1 else "Discovery Saja"
                row[header_name] = val if val is not None else ""
            rows.append(row)
        return rows

    @classmethod
    def generate_csv(cls, items: List[Dict[str, Any]]) -> bytes:
        flat_data = cls._prepare_flat_data(items)
        df = pd.DataFrame(flat_data)
        output = io.StringIO()
        # Tulis dengan BOM utf-8-sig agar karakter terbuka sempurna di Microsoft Excel Windows
        df.to_csv(output, index=False, encoding="utf-8-sig")
        return output.getvalue().encode("utf-8-sig")

    @classmethod
    def generate_excel(cls, items: List[Dict[str, Any]]) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Intelligence Kost"
        ws.views.sheetView[0].showGridLines = True

        # Header styling
        headers = [h for _, h in cls.COLUMN_MAPPING]
        ws.append(headers)

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border_thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align
            cell.border = border_thin
        ws.row_dimensions[1].height = 28

        # Kolom indeks kunci (1-indexed)
        col_price = headers.index("Harga Bulanan (Rp)") + 1
        col_vfm = headers.index("Skor VFM (0-100)") + 1
        col_motor_dist = headers.index("Jarak Motor (km)") + 1
        col_car_dist = headers.index("Jarak Mobil (km)") + 1
        col_walk_dist = headers.index("Jarak Jalan (km)") + 1
        col_straight = headers.index("Garis Lurus (km)") + 1
        col_motor_time = headers.index("Waktu Motor (mnt)") + 1
        col_car_time = headers.index("Waktu Mobil (mnt)") + 1
        col_walk_time = headers.index("Waktu Jalan (mnt)") + 1
        col_link_mami = headers.index("Link Mamikos") + 1
        col_link_gmaps = headers.index("Rute Google Maps") + 1

        # Data rows
        link_font = Font(name="Calibri", size=10, color="2563EB", underline="single")
        data_font = Font(name="Calibri", size=10)
        data_align_left = Alignment(horizontal="left", vertical="center")
        data_align_center = Alignment(horizontal="center", vertical="center")
        data_align_right = Alignment(horizontal="right", vertical="center")

        row_num = 2
        for item in items:
            row_vals = []
            for key, header_name in cls.COLUMN_MAPPING:
                val = item.get(key)
                if key == "facilities_room":
                    if isinstance(val, list):
                        val = ", ".join(str(f) for f in val)
                    elif isinstance(val, str) and val.startswith("["):
                        try:
                            val = ", ".join(str(f) for f in json.loads(val))
                        except Exception:
                            pass
                elif key == "is_detail_enriched":
                    val = "Lengkap (Enriched)" if val == 1 else "Discovery Saja"
                row_vals.append(val if val is not None else "")
            ws.append(row_vals)

            # Number formats and styling
            ws.row_dimensions[row_num].height = 20
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_num, column=c_idx)
                c.font = data_font
                c.border = border_thin
                c.alignment = data_align_center

            # Kolom nama dan alamat left-align
            ws.cell(row=row_num, column=2).alignment = data_align_left
            ws.cell(row=row_num, column=headers.index("Alamat Lengkap") + 1).alignment = data_align_left
            ws.cell(row=row_num, column=headers.index("Fasilitas Kamar") + 1).alignment = data_align_left

            # Format harga: Rp #,##0
            c_p = ws.cell(row=row_num, column=col_price)
            c_p.number_format = '"Rp "#,##0'
            c_p.alignment = data_align_right

            # Format jarak: 0.0 "km"
            for c_dist in [col_motor_dist, col_car_dist, col_walk_dist, col_straight]:
                cell_d = ws.cell(row=row_num, column=c_dist)
                cell_d.number_format = '0.0 "km"'
                cell_d.alignment = data_align_right

            # Format durasi: 0.0 "mnt"
            for c_dur in [col_motor_time, col_car_time, col_walk_time]:
                cell_t = ws.cell(row=row_num, column=c_dur)
                cell_t.number_format = '0.0 "mnt"'
                cell_t.alignment = data_align_right

            # Hyperlinks aktif
            url_mami = item.get("url_detail")
            if url_mami:
                c_mami = ws.cell(row=row_num, column=col_link_mami)
                c_mami.value = f'=HYPERLINK("{url_mami}", "Buka Mamikos")'
                c_mami.font = link_font

            url_maps = item.get("url_gmaps_route")
            if url_maps:
                c_maps = ws.cell(row=row_num, column=col_link_gmaps)
                c_maps.value = f'=HYPERLINK("{url_maps}", "Navigasi Maps")'
                c_maps.font = link_font

            row_num += 1

        total_rows = len(items)
        if total_rows > 0:
            last_row = total_rows + 1
            # Conditional formatting: 3-Color Scale pada Skor VFM (Merah -> Kuning -> Hijau)
            vfm_col_letter = get_column_letter(col_vfm)
            rule_vfm = ColorScaleRule(
                start_type="num", start_value=0, start_color="FCA5A5",   # Merah muda
                mid_type="num", mid_value=50, mid_color="FDE047",      # Kuning muda
                end_type="num", end_value=100, end_color="86EFAC"      # Hijau muda
            )
            ws.conditional_formatting.add(f"{vfm_col_letter}2:{vfm_col_letter}{last_row}", rule_vfm)

            # Conditional formatting: 3-Color Scale pada Jarak Motor (Hijau dekat -> Kuning -> Merah jauh)
            dist_col_letter = get_column_letter(col_motor_dist)
            rule_dist = ColorScaleRule(
                start_type="min", start_color="86EFAC",                 # Hijau (paling dekat)
                mid_type="percentile", mid_value=50, mid_color="FDE047", # Kuning
                end_type="max", end_color="FCA5A5"                      # Merah (paling jauh)
            )
            ws.conditional_formatting.add(f"{dist_col_letter}2:{dist_col_letter}{last_row}", rule_dist)

        # Freeze Panes: freeze row 1 dan column A-B (sehingga cell C2 menjadi freeze point)
        ws.freeze_panes = "C2"

        # Auto-filter
        ws.auto_filter.ref = ws.dimensions

        # Auto Column Width (min 12, max 45)
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = 0
            for cell in col:
                val_str = str(cell.value or "")
                if val_str.startswith("="):
                    val_str = "Buka Link Navigasi"
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(12, min(max_len + 3, 45))

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()
